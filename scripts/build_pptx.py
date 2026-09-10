"""Generate docs/slides/presentation.pptx from the same content as
docs/slides/presentation.html — real numbers from docs/problem_statement.md,
not re-derived. Build tool, not part of the app/runtime — needs
`pip install python-pptx` (not in requirements.txt, deliberately not a
runtime dependency). Re-run after editing this file to regenerate the
.pptx: `python scripts/build_pptx.py`.
"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- palette (light theme, matches docs/slides/presentation.html) ----
BG = RGBColor(0xF3, 0xF5, 0xEE)
SURFACE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1E, 0x2A, 0x20)
INK_SOFT = RGBColor(0x56, 0x63, 0x4F)
LINE = RGBColor(0xD7, 0xDE, 0xCB)
PASS = RGBColor(0x2E, 0x7D, 0x52)
PASS_SOFT = RGBColor(0xE4, 0xF0, 0xE4)
ISSUE = RGBColor(0xC4, 0x48, 0x2E)
ISSUE_SOFT = RGBColor(0xF7, 0xE4, 0xDD)

FONT_DISPLAY = "Georgia"  # Fraunces not available in PowerPoint -> closest serif fallback
FONT_BODY = "Calibri"
FONT_MONO = "Consolas"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def new_slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    return s


def textbox(slide, l, t, w, h):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tb.text_frame.word_wrap = True
    return tb


def set_run(run, text, size, color, font=FONT_BODY, bold=False, italic=False, spacing=None):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = font
    run.font.bold = bold
    run.font.italic = italic


def eyebrow(slide, text, top=Inches(0.5)):
    tb = textbox(slide, Inches(0.7), top, Inches(11.9), Inches(0.45))
    p = tb.text_frame.paragraphs[0]
    r1 = p.add_run()
    set_run(r1, "[ ", 13, PASS, FONT_MONO, bold=True)
    r2 = p.add_run()
    set_run(r2, text, 13, INK_SOFT, FONT_MONO, bold=False)
    r3 = p.add_run()
    set_run(r3, " ]", 13, PASS, FONT_MONO, bold=True)
    return tb


def heading(slide, text, top=Inches(1.05), width=Inches(11.9), size=32):
    tb = textbox(slide, Inches(0.7), top, width, Inches(1.5))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run()
    set_run(r, text, size, INK, FONT_DISPLAY, bold=True)
    return tb


def bullet_list(slide, items, left=Inches(0.7), top=Inches(2.2), width=Inches(6.6), size=15):
    tb = textbox(slide, left, top, width, Inches(4.5))
    tf = tb.text_frame
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(14)
        r = p.add_run()
        set_run(r, "—  " + item, size, INK_SOFT, FONT_BODY)
    return tb


def tag_card(slide, left, top, width, height, tag, tag_color, tag_bg, title, desc):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = SURFACE
    box.line.color.rgb = LINE
    box.line.width = Pt(1)
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.22)

    tagbox = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left + Inches(0.22), top + Inches(0.2), Inches(1.5), Inches(0.32))
    tagbox.fill.solid()
    tagbox.fill.fore_color.rgb = tag_bg
    tagbox.line.fill.background()
    tagbox.shadow.inherit = False
    tp = tagbox.text_frame
    tp.margin_left = Inches(0.08)
    tp.margin_top = 0
    tp.margin_bottom = 0
    trun = tp.paragraphs[0].add_run()
    set_run(trun, tag.upper(), 10, tag_color, FONT_MONO, bold=True)

    p1 = tf.paragraphs[0]
    p1.text = ""
    for _ in range(2):
        tf.add_paragraph()
    p_title = tf.paragraphs[1]
    r_title = p_title.add_run()
    set_run(r_title, title, 18, INK, FONT_DISPLAY, bold=True)
    p_desc = tf.paragraphs[2]
    r_desc = p_desc.add_run()
    set_run(r_desc, desc, 12, INK_SOFT, FONT_BODY)
    return box


def stat_tile(slide, left, top, width, height, big, cap, issue=False):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = SURFACE
    box.line.color.rgb = LINE
    box.line.width = Pt(1)
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_top = Inches(0.18)
    p1 = tf.paragraphs[0]
    r1 = p1.add_run()
    set_run(r1, big, 34, ISSUE if issue else PASS, FONT_DISPLAY, bold=True)
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    set_run(r2, cap, 11, INK_SOFT, FONT_MONO)
    return box


def callout(slide, left, top, width, height, label, text):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = ISSUE_SOFT
    box.line.color.rgb = ISSUE
    box.line.width = Pt(1.5)
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_top = Inches(0.2)
    tf.margin_right = Inches(0.2)
    p1 = tf.paragraphs[0]
    r1 = p1.add_run()
    set_run(r1, label.upper(), 10, ISSUE, FONT_MONO, bold=True)
    p2 = tf.add_paragraph()
    p2.space_before = Pt(6)
    r2 = p2.add_run()
    set_run(r2, text, 12.5, INK, FONT_BODY)
    return box


def lesson_row(slide, left, top, width, tag, text):
    tagbox = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.65), Inches(0.3))
    tagbox.fill.background()
    tagbox.line.color.rgb = PASS
    tagbox.line.width = Pt(1)
    tagbox.shadow.inherit = False
    tp = tagbox.text_frame
    tp.margin_left = Inches(0.05)
    tp.margin_top = 0
    tr = tp.paragraphs[0].add_run()
    set_run(tr, tag, 9, PASS, FONT_MONO, bold=True)

    tb = textbox(slide, left + Inches(0.85), top - Inches(0.03), width - Inches(0.85), Inches(0.5))
    r = tb.text_frame.paragraphs[0].add_run()
    set_run(r, text, 12.5, INK_SOFT, FONT_BODY)
    return tb


def table(slide, left, top, width, height, headers, rows, win_rows=()):
    n_rows = len(rows) + 1
    n_cols = len(headers)
    gt = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    t = gt.table
    for c, htext in enumerate(headers):
        cell = t.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = BG
        p = cell.text_frame.paragraphs[0]
        r = p.add_run()
        set_run(r, htext.upper(), 11, INK_SOFT, FONT_MONO, bold=True)
    for ri, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = t.cell(ri, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = PASS_SOFT if ri in win_rows else SURFACE
            p = cell.text_frame.paragraphs[0]
            r = p.add_run()
            color = PASS if ri in win_rows else INK
            set_run(r, str(val), 13, color, FONT_BODY, bold=(ri in win_rows))
    return gt


def chip(slide, left, top, text, size=14):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.2), Inches(0.5))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0xEA, 0xEE, 0xE1)
    box.line.color.rgb = LINE
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = False
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.08)
    p = tf.paragraphs[0]
    r = p.add_run()
    set_run(r, text, size, INK, FONT_MONO, bold=True)
    box.width = Inches(0.35 + len(text) * 0.095)
    return box


# ============================================================ Slide 0 — Cover
s = new_slide()
eyebrow(s, "COMPUTER VISION MODULE · AI ENGINEER K08-0226", top=Inches(1.6))
tb = textbox(s, Inches(0.7), Inches(2.15), Inches(11), Inches(1.8))
r = tb.text_frame.paragraphs[0].add_run()
set_run(r, "Yoga Pose Detection & Form Scoring", 44, INK, FONT_DISPLAY, bold=True)
tb2 = textbox(s, Inches(0.7), Inches(3.55), Inches(9.5), Inches(1.2))
r2 = tb2.text_frame.paragraphs[0].add_run()
set_run(r2, "Phát hiện + phân loại tư thế yoga bằng YOLOv8 tự train, chấm điểm form "
            "bằng góc khớp MediaPipe, deploy public thật với HTTPS và camera trên iPhone.",
        16, INK_SOFT, FONT_BODY)
chip(s, Inches(0.7), Inches(4.9), "demo · computer-vision-project.pthieu290998.workers.dev", size=12)

# ============================================================ Slide 1 — Bài toán
s = new_slide()
eyebrow(s, "MỞ ĐẦU · BÀI TOÁN")
heading(s, "Một model train được, làm cả định vị lẫn phân loại", width=Inches(7.2), size=28)
bullet_list(s, [
    "Input: ảnh/video học viên đang tập yoga",
    "Output: vị trí + tên tư thế, đánh giá form đúng/sai, gợi ý sửa",
    "Ràng buộc đề bài: bắt buộc 1 model classification/detection/segmentation tự train",
    "Lựa chọn: YOLOv8 detection — vừa localize vừa classify trong 1 bước, hỗ trợ nhiều học viên/khung hình cùng lúc",
], top=Inches(2.3), width=Inches(7.2))
stat_tile(s, Inches(8.3), Inches(2.3), Inches(4.3), Inches(2.0), "5", "LỚP TƯ THẾ\nbridge · downward · plank\nshoulderstand · tree")

# ============================================================ Slide 2 — Kiến trúc
s = new_slide()
eyebrow(s, "KIẾN TRÚC · 2 LỚP ĐỘC LẬP")
heading(s, "Model train được, model mượn, và luật tay — không trộn lẫn", width=Inches(11.9), size=26)
cw = Inches(3.9)
gap = Inches(0.25)
tag_card(s, Inches(0.7), Inches(2.6), cw, Inches(3.2), "trained", PASS, PASS_SOFT,
         "YOLOv8 (ONNX)", "Detect + classify tư thế trong 1 bước. Rubric mục 1-4, 7 chấm điểm lớp này.")
tag_card(s, Inches(0.7) + cw + gap, Inches(2.6), cw, Inches(3.2), "pretrained", RGBColor(0x8A, 0x6D, 0x1F), RGBColor(0xF5, 0xEC, 0xD2),
         "MediaPipe Pose", "33 keypoint mỗi box, dùng nguyên off-the-shelf, không train lại, không tính accuracy riêng.")
tag_card(s, Inches(0.7) + 2 * (cw + gap), Inches(2.6), cw, Inches(3.2), "rule-based", RGBColor(0x5A, 0x4A, 0x9C), RGBColor(0xEA, 0xE4, 0xF7),
         "score_pose()", "So góc khớp với POSE_RULES hiệu chỉnh thủ công. Rubric mục 6 — \u201cý tưởng riêng\u201d.")

# ============================================================ Slide 3 — Dữ liệu
s = new_slide()
eyebrow(s, "GIAI ĐOẠN 01 · DỮ LIỆU")
heading(s, "v1 chạy hết vòng đời thật — v2 dừng ở thiết kế", width=Inches(7.2), size=27)
bullet_list(s, [
    "v1: 5 lớp, format YOLO có sẵn bbox, nguồn Roboflow \u201cYOLO YOGA Dataset\u201d",
    "Augmentation: fliplr bật, flipud tắt (lật dọc phá vỡ ý nghĩa tư thế), degrees giới hạn",
    "v2 (15-20 lớp, Yoga-82 + pseudo-label): code đã thiết kế sẵn (bootstrap_bboxes()) nhưng không triển khai — v1 đã đủ robust (xem GĐ 05)",
], top=Inches(2.35), width=Inches(7.2))
stat_tile(s, Inches(8.3), Inches(2.35), Inches(4.3), Inches(1.15), "v1 · LIVE", "download → train → ablation → eval → export → deploy")
stat_tile(s, Inches(8.3), Inches(3.65), Inches(4.3), Inches(1.15), "v2 · STUB", "bootstrap_bbox.py — NotImplementedError, chưa cần chạy", issue=True)

# ============================================================ Slide 4 — Training & Ablation
s = new_slide()
eyebrow(s, "GIAI ĐOẠN 02-03 · TRAINING & ABLATION")
heading(s, "Cùng seed, đổi đúng 1 biến", size=30)
tb = textbox(s, Inches(0.7), Inches(1.95), Inches(10), Inches(0.4))
r = tb.text_frame.paragraphs[0].add_run()
set_run(r, "model yolov8n.pt  ·  seed 42  ·  epochs 50  ·  imgsz 640", 13, INK_SOFT, FONT_MONO)
table(s, Inches(0.7), Inches(2.5), Inches(10.5), Inches(1.6),
      ["Run", "Augmentation", "Kết quả"],
      [["Baseline", "ON", "Thắng — chọn làm config chính thức"],
       ["Ablation", "OFF", "So sánh, không chọn"]],
      win_rows={1})
tb2 = textbox(s, Inches(0.7), Inches(4.4), Inches(10.5), Inches(0.5))
r2 = tb2.text_frame.paragraphs[0].add_run()
set_run(r2, "Số mAP@0.5 / mAP@0.5:0.95 thật — xem bảng đầy đủ trong docs/problem_statement.md.", 12, INK_SOFT, FONT_BODY, italic=True)

# ============================================================ Slide 5 — Evaluation
s = new_slide()
eyebrow(s, "GIAI ĐOẠN 04 · EVALUATION & ERROR ANALYSIS")
heading(s, "Confusion matrix sạch, EigenCAM xác nhận nhìn đúng chỗ", width=Inches(11.5), size=27)
bullet_list(s, [
    "Confusion matrix: không có cặp lớp nào bị nhầm đáng kể trong tập test (cùng phân bố training)",
    "EigenCAM (3-5 ảnh): model tập trung đúng vùng cơ thể, không bị phân tâm bởi nền",
    "→ Không có lỗi trong tập test cùng phân bố để \u201csửa\u201d — phải đổi hướng cho bước feedback loop",
], top=Inches(2.4), width=Inches(11))

# ============================================================ Slide 6 — Feedback loop
s = new_slide()
eyebrow(s, "GIAI ĐOẠN 05 · FEEDBACK LOOP")
heading(s, "Đo lỗi thật ngoài phân bố, không sửa mù quáng", width=Inches(7.2), size=27)
bullet_list(s, [
    "18 ảnh thật, nguồn Kaggle — khác hoàn toàn Roboflow (khác chụp/nền/2 lớp)",
    "1 lỗi duy nhất, khả năng do nhãn gốc dataset sai — không phải lỗi hệ thống",
    "Kết luận: baseline đủ robust → không retrain",
], top=Inches(2.35), width=Inches(7.2))
stat_tile(s, Inches(8.3), Inches(2.35), Inches(4.3), Inches(1.8), "94.4%", "17/18 ẢNH OOD ĐÚNG")

# ============================================================ Slide 7 — Form scoring
s = new_slide()
eyebrow(s, "GIAI ĐOẠN 06 · Ý TƯỞNG RIÊNG")
heading(s, "Góc khớp thật hiệu chỉnh ngưỡng, không đoán số", width=Inches(11.5), size=27)
stat_tile(s, Inches(0.7), Inches(2.4), Inches(3.3), Inches(1.5), "14/15", "ẢNH MẪU THẬT ĐÚNG FORM")
stat_tile(s, Inches(0.7), Inches(4.05), Inches(3.3), Inches(1.5), "5/5", "CASE LANDMARK TỔNG HỢP — PHÁT HIỆN ĐÚNG")
callout(s, Inches(4.3), Inches(2.4), Inches(8.3), Inches(3.15), "2 hạn chế ghi nhận công khai",
        "MediaPipe kém tin cậy trên tư thế lộn ngược, nền tương phản thấp — landmark detect sai. "
        "Góc chiếu 2D bị méo khi camera chụp xiên — hạn chế cố hữu của ảnh đơn không có depth, không phải bug.")

# ============================================================ Slide 8 — Export & Backend
s = new_slide()
eyebrow(s, "GIAI ĐOẠN 07 · EXPORT & BACKEND")
heading(s, "ONNX nhanh hơn PyTorch ~17% — nhưng CPU free-tier là CPU free-tier", width=Inches(11.8), size=24)
table(s, Inches(0.7), Inches(2.5), Inches(10.5), Inches(2.6),
      ["Nơi đo", "Latency"],
      [["PyTorch · GPU T4 (Colab)", "247.94 ms"],
       ["ONNX Runtime · GPU T4 (Colab)", "206.29 ms"],
       ["ONNX Runtime · CPU (local)", "245.8 ms"],
       ["ONNX Runtime · CPU (Render free)", "~4.5 – 8.4 s"]],
      win_rows={2})
tb2 = textbox(s, Inches(0.7), Inches(5.35), Inches(10.5), Inches(0.6))
r2 = tb2.text_frame.paragraphs[0].add_run()
set_run(r2, "Chênh lệch CPU free-tier do tài nguyên bị giới hạn/chia sẻ mạnh — kết quả vẫn đúng, chỉ chậm hơn, không phải bug.", 12, INK_SOFT, FONT_BODY, italic=True)

# ============================================================ Slide 9 — Web demo & Deploy
s = new_slide()
eyebrow(s, "GIAI ĐOẠN 08-09 · WEB DEMO & DEPLOY")
heading(s, "2 origin thật, HTTPS thật, camera thật trên iPhone", width=Inches(11.8), size=26)
tb = textbox(s, Inches(0.7), Inches(2.35), Inches(5.4), Inches(2.5))
r = tb.text_frame.paragraphs[0].add_run()
set_run(r, "iPhone Safari  →  Cloudflare Pages (frontend tĩnh)  →  Render "
            "(backend Docker, free) — CORS qua env ALLOWED_ORIGINS.",
        14, INK_SOFT, FONT_BODY)
tbl = textbox(s, Inches(6.5), Inches(2.15), Inches(6.1), Inches(0.4))
rl = tbl.text_frame.paragraphs[0].add_run()
set_run(rl, "3 BUG MÔI TRƯỜNG THẬT, ĐÃ SỬA", 12, INK_SOFT, FONT_MONO, bold=True)
lesson_row(s, Inches(6.5), Inches(2.65), Inches(6.1), "FIX", "Render chọn nhầm runtime Python thay vì Docker → chọn đúng Environment")
lesson_row(s, Inches(6.5), Inches(3.25), Inches(6.1), "FIX", "Thiếu libGL.so.1 trên container Linux tối giản → thêm apt package")
lesson_row(s, Inches(6.5), Inches(3.85), Inches(6.1), "FIX", "iOS Safari đẩy camera vào native fullscreen player → webkit-playsinline")

# ============================================================ Slide 10 — Lỗi & bài học
s = new_slide()
eyebrow(s, "LỖI & BÀI HỌC")
heading(s, "Chẩn đoán từ log/screenshot thật — không đoán", width=Inches(11.8), size=27)
lesson_row(s, Inches(0.7), Inches(2.55), Inches(11.5), "FIX", "Tiền tố \u201cyoga-pose \u201d chưa bỏ trước khi so sánh → OOD báo nhầm 0% — sửa bằng đối chiếu log gốc.")
lesson_row(s, Inches(0.7), Inches(3.2), Inches(11.5), "FIX", "mediapipe \u2265 1.0 bỏ hẳn solutions.pose (API cũ) → pin đúng bản còn API, xác nhận qua lỗi build thật.")
lesson_row(s, Inches(0.7), Inches(3.85), Inches(11.5), "FIX", "Letterbox rounding lệch ~2.3px trên ảnh tỉ lệ cực đoan → làm tròn pad 1 lần, dùng lại nhất quán.")
lesson_row(s, Inches(0.7), Inches(4.5), Inches(11.5), "FIX", "3 bug deploy (runtime Docker, libGL, iOS fullscreen) — chẩn đoán trực tiếp từ log/screenshot thật.")

# ============================================================ Slide 11 — Demo
s = new_slide()
eyebrow(s, "TRÌNH DIỄN TRỰC TIẾP", top=Inches(2.6))
tb = textbox(s, Inches(0.7), Inches(3.1), Inches(11.9), Inches(1.0))
p = tb.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.LEFT
r = p.add_run()
set_run(r, "Camera thật · HTTPS thật · kết quả thật", 32, INK, FONT_DISPLAY, bold=True)
chip(s, Inches(0.7), Inches(4.1), "computer-vision-project.pthieu290998.workers.dev", size=14)
tb2 = textbox(s, Inches(0.7), Inches(4.9), Inches(10.5), Inches(0.8))
r2 = tb2.text_frame.paragraphs[0].add_run()
set_run(r2, "Quay trực tiếp lúc trình bày — camera trên điện thoại thật qua URL public, "
            "không dùng ảnh chụp màn hình dựng sẵn.", 14, INK_SOFT, FONT_BODY)

# ============================================================ Slide 12 — Kết
s = new_slide()
eyebrow(s, "KẾT")
heading(s, "Cảm ơn — câu hỏi?", size=32)
bullet_list(s, [
    "Repo: github.com/tom-the-prgrmr/computer-vision-project",
    "Demo live: computer-vision-project.pthieu290998.workers.dev",
    "Kiến trúc: docs/architecture.md",
    "Chi tiết 7 mục rubric: docs/problem_statement.md",
], top=Inches(2.3), width=Inches(11))

OUT_PATH = Path(__file__).resolve().parent.parent / "docs" / "slides" / "presentation.pptx"
prs.save(str(OUT_PATH))
print(f"Saved {len(prs.slides._sldIdLst)} slides -> {OUT_PATH}")
